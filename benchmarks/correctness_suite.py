"""Version-agnostic correctness checks for the AutoCAD MCP ezdxf backend.

Driven by ``compare_versions.py`` (the A/B runner). Run ONE check in isolation:

    python benchmarks/correctness_suite.py <check_name>
    python benchmarks/correctness_suite.py --list

Prints 'PASS' or 'FAIL' on the last line; a crash (e.g. SIGSEGV) or an
unsupported method => non-zero exit, which the runner counts as a miss.

The same file is executed against both the baseline worktree and the current
checkout (PYTHONPATH + cwd point at each repo), so a check that exercises a
method the old version lacks simply raises AttributeError -> miss.
"""

from __future__ import annotations

import asyncio
import math
import os
import sys
import tempfile


async def _b():
    from backends.ezdxf_backend import EzdxfBackend

    be = EzdxfBackend()
    await be.connect()
    await be.drawing_new()
    return be


# ── Core capability (expected to pass on BOTH versions — sanity / not rigged) ──


async def core_line_length():
    b = await _b()
    ln = await b.entity_create_line(0, 0, 3, 4)
    info = await b.entity_get(ln.handle)
    return abs(float(info.properties["length"]) - 5.0) < 1e-6


async def core_circle_radius():
    b = await _b()
    c = await b.entity_create_circle(0, 0, 10)
    info = await b.entity_get(c.handle)
    return abs(float(info.properties["radius"]) - 10.0) < 1e-6


async def core_layer_create():
    b = await _b()
    await b.layer_create("BENCH", color=1)
    names = {lyr.name for lyr in await b.layer_list()}
    return "BENCH" in names


async def core_polyline_closed():
    b = await _b()
    p = await b.entity_create_polyline([[0, 0], [10, 0], [10, 10]], closed=True)
    info = await b.entity_get(p.handle)
    return bool(info.properties.get("closed"))


async def core_linear_dim():
    b = await _b()
    d = await b.dimension_linear(0, 0, 100, 0, 50, 20)
    info = await b.entity_get(d.handle)
    return "DIM" in info.type.upper()


async def core_save_dxf_roundtrip():
    b = await _b()
    await b.entity_create_line(0, 0, 10, 10)
    path = os.path.join(tempfile.mkdtemp(), "rt.dxf")
    await b.drawing_save_as(path, "dxf")
    import ezdxf

    ezdxf.readfile(path)  # raises if not a valid DXF
    return os.path.getsize(path) > 0


# ── Dimensions (I16) ──────────────────────────────────────────────────────────


async def dim_aligned_no_error():
    b = await _b()
    d = await b.dimension_aligned(0, 0, 50, 50, 30, 70)
    return d is not None


async def dim_angular_no_error():
    b = await _b()
    d = await b.dimension_angular(0, 0, 100, 0, 0, 100, 60, 60)
    return d is not None


# ── Polar array full circle (R9) ──────────────────────────────────────────────


async def array_polar_360_distinct():
    b = await _b()
    c = await b.entity_create_circle(10, 0, 1)
    res = await b.entity_array_polar(c.handle, 4, 360.0, 0, 0)
    # 3 copies + original = 4 distinct centers, none coincident with the original.
    centers = {(10.0, 0.0)}
    for e in res:
        info = await b.entity_get(e.handle if hasattr(e, "handle") else e)
        cx, cy = info.properties["center"]
        centers.add((round(cx, 3), round(cy, 3)))
    return len(centers) == 4


# ── Deterministic geometry (I9) ────────────────────────────────────────────────


async def point_intersection_line_line():
    b = await _b()
    l1 = await b.entity_create_line(0, 0, 10, 0)
    l2 = await b.entity_create_line(5, -5, 5, 5)
    pt = await b.point_intersection(l1.handle, l2.handle)
    return abs(pt[0] - 5.0) < 1e-6 and abs(pt[1] - 0.0) < 1e-6


async def point_tangent_external():
    b = await _b()
    c = await b.entity_create_circle(0, 0, 3)
    tx, ty = await b.point_tangent(c.handle, 5, 0, ref_y=10)
    # tangency: (T-F).(T-C) == 0 and |T-C| == r
    fx, fy = 5.0, 0.0
    dot = (tx - fx) * (tx - 0) + (ty - fy) * (ty - 0)
    return abs(dot) < 1e-4 and abs(math.hypot(tx, ty) - 3.0) < 1e-4


# ── Selection / property parity (N3, N5) ───────────────────────────────────────


async def arc_has_length():
    b = await _b()
    a = await b.entity_create_arc(0, 0, 10, 0, 90)
    info = await b.entity_get(a.handle)
    return abs(float(info.properties["length"]) - (10 * math.pi / 2)) < 1e-3


async def arc_select_by_length():
    b = await _b()
    await b.entity_create_arc(0, 0, 10, 0, 90)  # length ~15.7
    sel = await b.entity_select_smart({"type": "ARC", "length_range": [15, 16]})
    return len(sel) == 1


async def ezdxf_bounding_box():
    b = await _b()
    ln = await b.entity_create_line(0, 0, 30, 40)
    info = await b.entity_get(ln.handle)
    bb = info.properties.get("bounding_box")
    return isinstance(bb, dict) and set(bb) == {"min", "max"}


async def mtext_rotation_roundtrip():
    b = await _b()
    mt = await b.entity_create_mtext("r", 0, 0, width=50, height=2.5, rotation=30.0)
    info = await b.entity_get(mt.handle)
    return abs(float(info.properties.get("rotation", 0.0)) - 30.0) < 1e-6


# ── Rendering (screenshot Agg / SIGSEGV) ───────────────────────────────────────


async def screenshot_png():
    b = await _b()
    await b.entity_create_line(0, 0, 100, 100)
    png = await b.view_screenshot()
    return png is not None and png[:4] == b"\x89PNG"


# ── Quality gate / critique (R4, N4) ───────────────────────────────────────────


async def dim_overlap_critique_fires():
    b = await _b()
    await b.drawing_apply_iso_layers("mech")
    await b.dimension_linear(0, 0, 100, 0, 50, 20, layer="DIM")
    await b.dimension_linear(0, 0, 100, 0, 50, 20, layer="DIM")
    issues = await b.drawing_critique(focus=["dim_overlap"])
    return len(issues) >= 1


async def iso13567_dim_layer():
    b = await _b()
    await b.drawing_apply_iso_layers("iso13567")
    ln = await b.entity_create_line(0, 0, 80, 0, layer="M-GEOMET-E-N")
    dims = await b.dimension_auto([ln.handle], style="chain")
    return dims[0].layer == "M-DIMEN-T-N"


async def construction_left_iso_caught():
    b = await _b()
    await b.drawing_apply_iso_layers("iso13567")
    await b.construction_xline(0, 0, 90)
    issues = await b.drawing_critique(focus=["construction_left"])
    return len(issues) == 1


# ── Engineering (gear) ──────────────────────────────────────────────────────────


async def gear_no_self_overlap():
    from engineering.gear import generate_full_gear_outline

    pts = generate_full_gear_outline(2, 50, 20.0)  # z=50 -> root_r >= base_r
    root_r = (2 * 50 / 2.0) - 1.25 * 2
    return min(math.hypot(x, y) for x, y in pts) >= root_r - 1e-6


# ── Linetype on demand (R24) ───────────────────────────────────────────────────


async def center_linetype_applied():
    b = await _b()
    ln = await b.entity_create_line(0, 0, 50, 0, linetype="CENTER")
    info = await b.entity_get(ln.handle)
    return info.linetype.upper() == "CENTER"


# ── Measurement (v1.5.0) ───────────────────────────────────────────────────────
#
# These three are NEW capability, not repaired regressions: v1.4.0 has no
# `entity_measure` and no boundary tracing, so it misses them by not having the
# method rather than by getting it wrong. They are here because all three broke
# during v1.5.0 development and nothing outside their own unit tests would have
# noticed.


async def hatch_area_subtracts_its_island():
    """A section view is mostly holes; counting them is a 33% error here."""
    b = await _b()
    h = await b.entity_create_hatch("ANSI31", [[0, 0], [20, 0], [20, 20], [0, 20]])
    await b.hatch_add_boundary(
        h.handle,
        [
            {"type": "line", "start": [5, 5], "end": [15, 5]},
            {"type": "line", "start": [15, 5], "end": [15, 15]},
            {"type": "line", "start": [15, 15], "end": [5, 15]},
            {"type": "line", "start": [5, 15], "end": [5, 5]},
        ],
    )
    return abs((await b.entity_measure(h.handle))["area"] - 300.0) < 1e-6


async def boundary_area_agrees_with_measure():
    """Two tools, one shape, one number."""
    b = await _b()
    await b.entity_create_line(0, 0, 10, 0)
    await b.entity_create_line(10, 0, 10, 10)
    await b.entity_create_line(0, 10, 0, 0)
    await b.entity_create_arc(5, 10, 5, 0, 180)
    loop = await b.boundary_trace(5.0, 5.0)
    measured = await b.entity_measure(loop["handle"])
    true_area = 100.0 + math.pi * 25.0 / 2.0
    return abs(loop["area"] - measured["area"]) < 1e-9 and abs(measured["area"] - true_area) < 1e-5


async def two_vertex_circle_has_area():
    """DXF's most compact circle. The shoelace over its two points is 0."""
    from engineering.measure import polygon_area_perimeter

    area, _ = polygon_area_perimeter([(-10.0, 0.0, 1.0), (10.0, 0.0, 1.0)], closed=True)
    return abs(area - math.pi * 100.0) < 1e-9


async def diameter_dim_measures_the_diameter():
    """The number on the callout is the whole point of the callout.

    `add_diameter_dim` measures to `mpoint`; the leader is a text placement.
    Feeding it centre + (radius + leader) made every diameter come back
    2 x leader_length too large at DEFAULT settings - 60 on a true 40.
    """
    b = await _b()
    info = await b.dimension_diameter(-20, 0, 20, 0)
    dim = b._doc.entitydb.get(info.handle)
    return abs(dim.get_measurement() - 40.0) < 1e-6


async def radius_dim_ignores_the_leader_length():
    b = await _b()
    near = await b.dimension_radius(0, 0, 20, 0, leader_length=5)
    far = await b.dimension_radius(0, 0, 20, 0, leader_length=40)
    values = [
        b._doc.entitydb.get(near.handle).get_measurement(),
        b._doc.entitydb.get(far.handle).get_measurement(),
    ]
    return all(abs(v - 20.0) < 1e-6 for v in values)


# ── P&ID (v1.6, track A) ──────────────────────────────────────────────────────


async def pid_block_define_attdef_roundtrip():
    b = await _b()
    await b.block_define(
        "PID_T",
        [{"type": "line", "x1": 0, "y1": 0, "x2": 8, "y2": 0}],
        [{"tag": "TAG", "x": 0, "y": 3, "height": 2.5}],
    )
    ref = await b.block_insert("PID_T", 5, 5, attributes={"TAG": "HV-1"})
    return await b.block_get_attributes(ref.handle) == {"TAG": "HV-1"}


async def pid_tag_parse_fic():
    from engineering.pid.tags import parse_tag

    return parse_tag("FIC-101")["description"] == "Flow Indicating Controller"


async def pid_graph_edge_count():
    from engineering.pid.spec import EXAMPLE_SPEC, run_spec

    b = await _b()
    result = await run_spec(b, EXAMPLE_SPEC)
    return len(result["graph"]["edges"]) == 4 and result["graph"]["stats"]["dangling"] == 0


# ── Settings and environment (v1.6, track E) ─────────────────────────────────


async def settings_dimstyle_iso25_values():
    """The ISO-25 preset reaches the DIMSTYLE table with ISO 129-1's numbers.

    Read back through ezdxf directly, not through `dimstyle_list`, so the
    check cannot be satisfied by a list that echoes its own input.
    """
    from engineering.standards.dimstyles import resolve_dimstyle

    b = await _b()
    await b.dimstyle_create("ISO-25", resolve_dimstyle("iso-25", None), set_current=True)
    style = b._doc.dimstyles.get("ISO-25")
    listed = {row["name"]: row for row in await b.dimstyle_list()}
    return (
        abs(float(style.dxf.dimtxt) - 2.5) < 1e-9
        and abs(float(style.dxf.dimasz) - 2.5) < 1e-9
        and abs(float(style.dxf.dimexe) - 1.25) < 1e-9
        and abs(float(style.dxf.dimexo) - 0.625) < 1e-9
        and abs(float(style.dxf.dimgap) - 0.625) < 1e-9
        and int(style.dxf.dimtad) == 1
        and int(style.dxf.dimdsep) == ord(",")
        and str(style.dxf.dimtxsty) == "ISOCP"
        and b._doc.header["$DIMSTYLE"] == "ISO-25"
        and listed["ISO-25"]["current"] is True
        and listed["ISO-25"]["values"]["DIMDSEP"] == ","
    )


async def settings_layer_state_roundtrip():
    """Save → change → restore puts the table back, and the state survives a
    save/reopen because it is an XRECORD in the file, not server memory."""
    b = await _b()
    await b.layer_create("LS_A", color=1)
    await b.layer_create("LS_B", color=2)
    saved = await b.layer_state_save("bench", description="A/B round trip")
    await b.layer_modify("LS_A", color=5)
    await b.layer_freeze("LS_B")
    restored = await b.layer_state_restore("bench")
    layers = {layer.name: layer for layer in await b.layer_list()}
    path = os.path.join(tempfile.mkdtemp(), "ls.dxf")
    await b.drawing_save_as(path, "dxf")
    await b.drawing_open(path)
    names_after_reopen = {row["name"] for row in await b.layer_state_list()}
    return (
        saved["ok"] is True
        and restored["missing_layers"] == []
        and restored["new_layers"] == []
        and layers["LS_A"].color == 1
        and layers["LS_B"].is_frozen is False
        and "bench" in names_after_reopen
    )


def _sheet_near(size_mm, width_mm: float, height_mm: float) -> bool:
    """Within the spec's ±0.5 mm: COM measures ANSI B as 431.8 × 279.4."""
    try:
        w, h = (float(v) for v in size_mm)
    except (TypeError, ValueError):
        return False
    return abs(w - width_mm) < 0.5 and abs(h - height_mm) < 0.5


async def settings_pdf_mediabox_a3():
    """Page setup evidence is the PDF's own /MediaBox (spec §8.3).

    A fresh document's Layout1 is already ISO A3 landscape, so an A3-only
    read-back would pass with the setter stubbed out. The sheet is therefore
    moved to ANSI B first — 432 × 279, a size no fresh document has (spec §9
    names it) — and that PDF is read; then ISO A3 is applied and read again.
    The A3 apply must also report the layout moved from B to A3, so the PDF
    and the setter's own diff have to agree on where the sheet was.
    """
    from engineering.standards.papers import resolve_page_setup
    from engineering.standards.plot import batch_plot

    b = await _b()
    await b.page_setup_apply("Layout1", resolve_page_setup("ANSI_B"))
    out_dir = tempfile.mkdtemp()
    report_b = await batch_plot(b, ["Layout1"], out_dir, "{drawing}-{layout}-ansi_b.pdf")
    width_b, height_b = report_b["sheets"][0]["mediabox_mm"]
    applied = await b.page_setup_apply("Layout1", resolve_page_setup("ISO_A3"))
    report_a3 = await batch_plot(b, ["Layout1"], out_dir, "{drawing}-{layout}-iso_a3.pdf")
    width_a3, height_a3 = report_a3["sheets"][0]["mediabox_mm"]
    moved = applied["changed"].get("size_mm") or [[0.0, 0.0], [0.0, 0.0]]
    return (
        abs(width_b - 432.0) < 0.5
        and abs(height_b - 279.0) < 0.5
        and _sheet_near(moved[0], 432.0, 279.0)
        and _sheet_near(moved[1], 420.0, 297.0)
        and abs(width_a3 - 420.0) < 0.5
        and abs(height_a3 - 297.0) < 0.5
    )


# ── Mechanical parts and the sheet (tracks B + G) ───────────────────────────


def _field(row: dict, *names):
    """A row's value under whichever of these names it carries.

    The parts-list column spelling belongs to the sheet module; what this
    suite pins is the numbers, so it reads by meaning rather than by name.
    """
    for name in names:
        if name in row:
            return row[name]
    return None


def _inserted_handle(result: dict):
    return result.get("handle") or (result.get("handles") or [None])[0]


async def mech_part_roundtrip():
    """Draw a part, read it back from its own ACADMCP_MECH XDATA, compare.

    Segments and material only: a feature's parameter spelling belongs to
    `engineering/mech/features.py` and its own tests, and a check that guessed
    at it would fail for the wrong reason.
    """
    from engineering.mech.draw import draw_part, read_part
    from engineering.mech.part import build_part, part_to_dict

    b = await _b()
    spec = {
        "kind": "revolved",
        "name": "SHAFT",
        "material": "steel",
        "segments": [
            {"length": 40.0, "d_outer": 30.0},
            {"length": 20.0, "d_outer": 50.0, "d_inner": 10.0},
        ],
    }
    part = build_part(spec)
    drawn = await draw_part(b, part, at=(0.0, 0.0), views=("front",), dimension=False)
    back = await read_part(b, drawn["part_id"])
    return part_to_dict(build_part(back["part"])) == part_to_dict(part)


async def mech_section_hatch_area():
    """The hatched cut area of a full section, measured against a hand number.

    A plain sleeve: length 60, outside diameter 40, bore 20. A full
    longitudinal section cuts two faces, each 60 x 10, so the hatch fills
    60 * (40 - 20) = 1200 mm^2 exactly. The number is read out of the drawing
    through entity_measure (the tool's analysis_measure_entity), never from the
    code that drew it — which is why it cannot be faked.
    """
    from engineering.mech.draw import add_view, draw_part
    from engineering.mech.part import build_part

    b = await _b()
    part = build_part(
        {
            "kind": "revolved",
            "name": "SLEEVE",
            "material": "steel",
            "segments": [{"length": 60.0, "d_outer": 40.0, "d_inner": 20.0}],
        }
    )
    drawn = await draw_part(b, part, at=(0.0, 0.0), views=("front",), dimension=False)
    await add_view(
        b,
        drawn["part_id"],
        "section",
        plane={"p1": [0.0, 0.0], "p2": [60.0, 0.0], "label": "A"},
        style="full",
    )
    total = 0.0
    for ent in await b.entity_list(type_filter="HATCH", limit=500):
        total += float((await b.entity_measure(ent.handle))["area"])
    return abs(total - 1200.0) < 0.01


async def mech_iso286_on_dimension():
    """A 40 H7 bore dimension carries ISO 286's +0.025 / 0 through the layout.

    ISO 286-1 table 1: the 30-50 mm step has IT7 = 25 um, and H fixes the lower
    deviation at 0. The fit has to survive `layout_dimensions`, or the number
    is right in the table and wrong on the sheet.
    """
    from engineering.fits import fit_lookup
    from engineering.mech.dimension import layout_dimensions
    from engineering.mech.primitives import Circle, DimIntent
    from engineering.mech.views import View

    deviation = fit_lookup("H7", 40.0)
    if abs(deviation.upper_mm - 0.025) > 1e-9 or abs(deviation.lower_mm) > 1e-9:
        return False
    intent = DimIntent("diameter", (-20.0, 0.0), (20.0, 0.0), fit="H7", feature="bore")
    view = View(
        kind="front",
        prims=(Circle((0.0, 0.0), 20.0),),
        dims=(intent,),
        omitted=(),
        bbox=(-20.0, -20.0, 20.0, 20.0),
    )
    placed = layout_dimensions(view, [intent])
    return len(placed) == 1 and placed[0].fit == "H7" and placed[0].feature == "bore"


async def mech_thread_unrepresented_is_caught():
    """The ISO 6410 focus fires on a thread drawn as a plain circle.

    A quality gate that never fires is not a gate; this is the mechanical twin
    of `dim_overlap_critique_fires`.
    """
    from engineering.mech.critique import build_index, issues_for
    from engineering.mech.xdata import APP_ID, to_values

    b = await _b()
    payload = {
        "v": 1,
        "kind": "part",
        "id": "P1",
        "view": "side",
        "part": {
            "name": "STUD",
            "material": "steel",
            "segments": [{"length": 60.0, "d_outer": 20.0}],
            "features": [{"kind": "thread", "id": "t1", "designation": "M20x1.5"}],
        },
    }
    anchor = await b.entity_create_circle(0.0, 0.0, 0.5, layer="0")
    await b.entity_set_xdata(anchor.handle, APP_ID, to_values(payload))
    await b.entity_create_circle(0.0, 0.0, 10.0, layer="GEOMETRY")
    before = issues_for("mech_thread_unrepresented", await build_index(b))
    await b.entity_create_arc(0.0, 0.0, 8.0, 0.0, 270.0, layer="GEOMETRY")
    after = issues_for("mech_thread_unrepresented", await build_index(b))
    return len(before) == 1 and after == []


async def sheet_frame_iso5457_a3():
    """ISO 5457 A3: 420 x 297 trimmed, a 20 mm filing margin on the left and
    10 mm on the other three edges, so the frame runs (20, 10) to (410, 287)
    and the drawing space is 390 x 277. The trimmed sheet edge is drawn too
    (ISO 5457 4.1), so the outermost extent is the sheet, 0,0 to 420,297 -
    and nothing may reach past it."""
    from engineering.mech.primitives import points_of
    from engineering.sheet.frames import SHEETS, frame_prims

    if tuple(float(v) for v in SHEETS["A3"]) != (297.0, 420.0):
        return False
    prims = frame_prims("A3", orientation="landscape", zones=False, marks=False)
    points = {
        (round(float(x), 6), round(float(y), 6)) for prim in prims for x, y in points_of(prim)
    }
    frame = {(20.0, 10.0), (410.0, 10.0), (410.0, 287.0), (20.0, 287.0)}
    sheet = {(0.0, 0.0), (420.0, 0.0), (420.0, 297.0), (0.0, 297.0)}
    xs = [x for x, _y in points]
    ys = [y for _x, y in points]
    extent = (min(xs), min(ys), max(xs), max(ys))
    return frame <= points and sheet <= points and extent == (0.0, 0.0, 420.0, 297.0)


async def bom_balloon_link():
    """Two identical bolts collapse into one row of quantity 2, the item number
    is stable across a rerun, and the balloon drawn for that row carries it."""
    from engineering.mech.stdparts import insert_std_part
    from engineering.mech.xdata import APP_ID, decode
    from engineering.sheet.bom import balloon_prims, rows_from_records

    b = await _b()
    records = []
    for x in (0.0, 60.0):
        result = await insert_std_part(b, "ISO 4014 - M12x60", at=(x, 0.0), view="side")
        raw = await b.entity_get_xdata(_inserted_handle(result), APP_ID)
        records.append(decode(raw["xdata"][APP_ID]))
    rows = rows_from_records(records)
    again = rows_from_records(records)
    if len(rows) != 1 or len(again) != 1:
        return False
    item = _field(rows[0], "item", "no", "pos")
    if item is None or item != _field(again[0], "item", "no", "pos"):
        return False
    if int(_field(rows[0], "qty", "quantity") or 0) != 2:
        return False
    texts = [
        p.text for p in balloon_prims(int(item), (0.0, 40.0), (0.0, 5.0)) if hasattr(p, "text")
    ]
    return texts == [str(int(item))]


async def std_part_iso4014_m12():
    """ISO 4014's M12 row and the block it becomes.

    ISO 4014:2011 table 1, M12: width across flats s = 18 mm, head height
    k = 7.5 mm. The row is matched by value rather than by column name — the
    spelling is the catalogue's business, the numbers are the standard's — and
    the inserted block has to carry its designation back in ACADMCP_MECH.
    The catalogue nests a row's sizes one level down (``dims``), so the
    values are read one level deep.
    """
    from engineering.mech.stdparts import catalogue, insert_std_part
    from engineering.mech.xdata import APP_ID, decode

    rows = [row for row in catalogue("M12") if any("4014" in str(value) for value in row.values())]
    if not rows:
        return False
    values = [
        inner
        for row in rows
        for value in row.values()
        for inner in (value.values() if isinstance(value, dict) else (value,))
    ]
    numbers = {
        round(float(value), 3)
        for value in values
        if isinstance(value, (int, float)) and not isinstance(value, bool)
    }
    if not {18.0, 7.5} <= numbers:
        return False
    b = await _b()
    result = await insert_std_part(b, "ISO 4014 - M12x60", at=(0.0, 0.0), view="side")
    raw = await b.entity_get_xdata(_inserted_handle(result), APP_ID)
    payload = decode(raw["xdata"][APP_ID])
    return payload["kind"] == "std_part" and "M12" in str(payload["designation"])


# ── Architecture (track F) ──────────────────────────────────────────────────


def _arch_points(geometry) -> set:
    """Every endpoint of a drawn face or jamb, rounded to a micron."""
    return {
        (round(float(x), 6), round(float(y), 6))
        for line in (*geometry.faces, *geometry.jambs)
        for x, y in (line.p1, line.p2)
    }


def _arch_passes_through(geometry, point) -> bool:
    """Does a drawn face run *through* ``point`` (not merely end on it)?"""
    px, py = point
    for line in geometry.faces:
        (x1, y1), (x2, y2) = line.p1, line.p2
        length = math.hypot(x2 - x1, y2 - y1)
        cross = (x2 - x1) * (py - y1) - (y2 - y1) * (px - x1)
        along = ((px - x1) * (x2 - x1) + (py - y1) * (y2 - y1)) / (length * length)
        if abs(cross) <= 1e-6 * length and 1e-9 < along < 1.0 - 1e-9:
            return True
    return False


def _arch_region_area(geometry) -> float:
    from engineering.measure import polygon_area_perimeter

    return sum(
        polygon_area_perimeter([(x, y, 0.0) for x, y in loop], True)[0] for loop in geometry.regions
    )


async def arch_junction_l_t_x():
    """The three junction kinds, against coordinates worked out by hand.

    L: one 200 mm centre wall turning at (4000, 0) - the outer corner is
    (4100, -100), the inner (3900, 100), and a mitred corner keeps the area at
    axis length x thickness: 7000 x 200 = 1 400 000 mm².
    T: a 100 mm stem on x = 3000 ending on the axis of a 200 mm wall - its faces
    stop on the near face at (2950, 100) and (3050, 100), the near face is open
    between them, the far face runs through; 6000 x 200 + 2900 x 100 = 1 490 000.
    X: two 200 mm walls crossing at (3000, 0) - the first stays whole, the second
    stops on its faces, all four faces are cut at x = 2900 / 3100 and y = +-100;
    6000 x 200 + 2 x 2900 x 200 = 2 360 000.
    """
    from engineering.arch.model import wall_from_dict
    from engineering.arch.walls import wall_geometry

    def wall(wall_id, axis, thickness):
        return wall_from_dict({"id": wall_id, "axis": axis, "thickness": thickness})

    l_geo = wall_geometry([wall("L", [[0, 0], [4000, 0], [4000, 3000]], 200)], poche=False)
    t_geo = wall_geometry(
        [wall("A", [[0, 0], [6000, 0]], 200), wall("B", [[3000, 0], [3000, 3000]], 100)],
        poche=False,
    )
    x_geo = wall_geometry(
        [wall("A", [[0, 0], [6000, 0]], 200), wall("B", [[3000, -3000], [3000, 3000]], 200)],
        poche=False,
    )
    l_ok = (
        {(4100.0, -100.0), (3900.0, 100.0)} <= _arch_points(l_geo)
        and not {(4100.0, 100.0), (3900.0, -100.0)} & _arch_points(l_geo)
        and abs(_arch_region_area(l_geo) - 1_400_000.0) < 1e-6
    )
    t_ok = (
        {(2950.0, 100.0), (3050.0, 100.0)} <= _arch_points(t_geo)
        and not _arch_passes_through(t_geo, (3000.0, 100.0))
        and _arch_passes_through(t_geo, (3000.0, -100.0))
        and abs(_arch_region_area(t_geo) - 1_490_000.0) < 1e-6
    )
    x_ok = (
        {(2900.0, 100.0), (3100.0, 100.0), (2900.0, -100.0), (3100.0, -100.0)}
        <= _arch_points(x_geo)
        and not any(
            _arch_passes_through(x_geo, p)
            for p in ((3000.0, 100.0), (3000.0, -100.0), (2900.0, 0.0), (3100.0, 0.0))
        )
        and abs(_arch_region_area(x_geo) - 2_360_000.0) < 1e-6
    )
    return l_ok and t_ok and x_ok


async def arch_room_area_net():
    """A 4 x 5 m room between 200 mm walls measures 20.00 m², not its axis area.

    The ring's axis runs (-100, -100) to (4100, 5100), so its inner faces are
    x = 0 / 4000 and y = 0 / 5000: the floor is 4000 x 5000 = 20 000 000 mm²,
    while the axis encloses 4200 x 5200 = 21 840 000. The number is read back
    out of the room record the label carries, and the label text must say it.
    """
    from engineering.arch.draw import draw_walls, read_plan
    from engineering.arch.lang import fmt_area_m2
    from engineering.arch.model import wall_from_dict
    from engineering.arch.rooms import label_room

    b = await _b()
    ring = wall_from_dict(
        {
            "id": "W1",
            "axis": [[-100, -100], [4100, -100], [4100, 5100], [-100, 5100]],
            "thickness": 200,
            "closed": True,
        }
    )
    await draw_walls(b, [ring])
    await label_room(b, {"id": "R1", "name": "ROOM", "number": "01", "at": [2000, 2500]})
    (room,) = (await read_plan(b))["rooms"]
    texts = [e.properties.get("text", "") for e in await b.entity_list(type_filter="TEXT")]
    return (
        abs(float(room.area) - 20_000_000.0) < 0.5
        and fmt_area_m2(room.area, "en") == "20.00 m²"
        and any("20.00 m²" in str(text) for text in texts)
    )


async def arch_opening_cuts_wall():
    """A 900 door at 1000 along a 5 m, 200 mm wall interrupts both faces.

    No face runs through x = 1450 (the middle of the door) on either face line
    y = +-100 while one still runs through x = 500, and two jambs close the cut
    at x = 1000 and x = 1900.
    """
    from engineering.arch.model import opening_from_dict, wall_from_dict
    from engineering.arch.walls import wall_geometry

    wall = wall_from_dict({"id": "W1", "axis": [[0, 0], [5000, 0]], "thickness": 200})
    door = opening_from_dict(
        {"id": "D1", "wall": "W1", "kind": "door", "offset": 1000, "width": 900}
    )
    geo = wall_geometry([wall], [door], poche=False)
    jambs = {
        round(float(line.p1[0]), 6)
        for line in geo.jambs
        if abs(line.p1[0] - line.p2[0]) < 1e-9 and abs(abs(line.p1[1] - line.p2[1]) - 200.0) < 1e-9
    }
    return (
        not _arch_passes_through(geo, (1450.0, 100.0))
        and not _arch_passes_through(geo, (1450.0, -100.0))
        and _arch_passes_through(geo, (500.0, 100.0))
        and {1000.0, 1900.0} <= jambs
        and tuple(geo.omitted) == ()
    )


async def arch_rooms_detect_foreign():
    """A plan of plain lines - no record of ours - yields the net areas.

    Two rooms drawn as the inner outlines of 200 mm walls on a layer named
    WALLS: 4000 x 5000 and 3000 x 5000, inside an outer outline (-200, -200)
    to (7400, 5200). The reader must find exactly those two faces - 20 000 000
    and 15 000 000 mm² - at the foreign-plan confidence 0.6, and must not have
    written anything.
    """
    from engineering.arch.rooms import rooms_detect

    b = await _b()
    for x0, y0, x1, y1 in (
        (-200.0, -200.0, 7400.0, 5200.0),
        (0.0, 0.0, 4000.0, 5000.0),
        (4200.0, 0.0, 7200.0, 5000.0),
    ):
        corners = [(x0, y0), (x1, y0), (x1, y1), (x0, y1)]
        for a, c in zip(corners, corners[1:] + corners[:1], strict=True):
            await b.entity_create_line(a[0], a[1], c[0], c[1], layer="WALLS")
    before = len(await b.entity_list(limit=1000))
    found = await rooms_detect(b)
    areas = sorted(round(float(room["area"]), 3) for room in found["rooms"])
    return (
        areas == [15_000_000.0, 20_000_000.0]
        and all(float(room["confidence"]) == 0.6 for room in found["rooms"])
        and len(await b.entity_list(limit=1000)) == before
    )


# ── Understanding foreign drawings (track H) ───────────────────────────────


def _plant_pair() -> dict:
    """The synthetic plant pair, written to a fresh temporary folder."""
    from tests.fixtures.plant_pair import build_plant_pair

    return build_plant_pair(tempfile.mkdtemp(prefix="acadmcp_plant_"))


def _quoted(handle: str) -> str:
    """A handle as it appears in a JSON dump - quoted, so '2A' never matches '12A3'."""
    return f'"{handle}"'


async def scale_check_detects_schematic():
    """The synthetic P&ID is schematic against its layout; a uniform 2x copy is not.

    The pair: room 1 of the P&ID is the layout stretched 1.3x, room 2 is
    rearranged, and of the 55 pairs of the 11 shared tags 21.8 % lie within
    ±10 % of the median ratio - at most 50 % is ``schematic`` (spec §6).
    The control: six tags at (0, 0), (4000, 0), (8000, 0), (0, 3000),
    (4000, 3000), (8000, 6000) and the same tags at twice those coordinates -
    the nearest pair is 3000 mm apart, every ratio is exactly 2.0, so the
    doubled drawing is ``to_scale`` against the first with factor 2.
    """
    import ezdxf

    from engineering.understand.scale import scale_check
    from engineering.understand.snapshot import read_snapshot

    truth = _plant_pair()
    pair = scale_check(read_snapshot(truth["pid"]), read_snapshot(truth["layout"]))
    folder = tempfile.mkdtemp(prefix="acadmcp_scale_")
    positions = [(0, 0), (4000, 0), (8000, 0), (0, 3000), (4000, 3000), (8000, 6000)]
    paths = []
    for factor in (1.0, 2.0):
        doc = ezdxf.new("R2018")
        msp = doc.modelspace()
        for number, (x, y) in enumerate(positions, start=1):
            msp.add_circle((x * factor, y * factor), 500.0)
            msp.add_text(f"T{number}", dxfattribs={"insert": (x * factor, y * factor)})
        path = os.path.join(folder, f"x{factor:g}.dxf")
        doc.saveas(path)
        paths.append(path)
    control = scale_check(read_snapshot(paths[1]), read_snapshot(paths[0]))
    return (
        pair["verdict"] == "schematic"
        and control["verdict"] == "to_scale"
        and abs(float(control["factor"]) - 2.0) < 1e-9
    )


async def pipe_takeoff_rmst_exact():
    """Every routable pipe run's layout length is the rectilinear MST of its tags.

    Measured on the synthetic pair with ``length_source="layout"``. The truth
    is the generator's own Manhattan / Prim arithmetic over the tag positions it
    placed, hand-checked in tests/test_plant_pair.py - PR1 5500 (5000 + 500),
    PR2 13200 (the three distances 6000 / 7200 / 13200, spanning tree 6000 +
    7200), CS1 20700 (20500 + 200). All eight routable runs must agree to a
    micron; the run ending on T105, which the layout lacks, is not routable.
    """
    from engineering.understand.network import build_network
    from engineering.understand.snapshot import read_snapshot
    from engineering.understand.takeoff import pipe_rows

    truth = _plant_pair()
    pid, layout = read_snapshot(truth["pid"]), read_snapshot(truth["layout"])
    services = tuple(sorted({run["service"] for run in truth["pipe_runs"]}))
    result = pipe_rows(build_network(pid), pid, layout, services=services, length_source="layout")
    found = {(r["service"], tuple(sorted(r["tags"]))): r for r in result["runs"]}
    routable = [run for run in truth["pipe_runs"] if run["layout_length_mm"] is not None]
    for run in routable:
        got = found.get((run["service"], tuple(run["tags"])))
        if got is None or got.get("layout_m") is None:
            return False
        if abs(float(got["layout_m"]) * 1000.0 - run["layout_length_mm"]) > 1e-6:
            return False
    return len(routable) == 8


async def cable_takeoff_roundup():
    """Cable rows: the panel from the callout, Manhattan length, metres rounded up.

    The truth rows of the synthetic pair, by hand: M11 (5500, 8200) to CP-1
    (1200, 10800) is 6900 mm, 6.9 m x 1.2 = 8.28 -> 9; T102 (8000, 3500) to CP-1
    is 14100 mm -> 16.92 -> 17; M12 (26300, 8400) to CP-2 (35000, 1300) is
    15800 mm -> 18.96 -> 19; T202 (28700, 3000) to CP-2 is 8000 mm -> 9.6 -> 10.
    The heater on T202 states no power, and its row must keep an empty one.
    """
    from engineering.understand.snapshot import read_snapshot
    from engineering.understand.takeoff import cable_rows

    truth = _plant_pair()
    rows = {
        row["tag"]: row
        for row in cable_rows(
            read_snapshot(truth["pid"]), read_snapshot(truth["layout"]), panel_rule="stated"
        )["rows"]
    }
    expected = {"M11": ("CP-1", 9), "T102": ("CP-1", 17), "M12": ("CP-2", 19), "T202": ("CP-2", 10)}
    for tag, (panel, metres) in expected.items():
        row = rows.get(tag)
        if row is None or row.get("panel") != panel or row.get("cable_m") != metres:
            return False
    return rows["T202"].get("kw") is None and rows["M11"].get("kw") == 5.5


async def diff_detects_known_edits():
    """Two revisions differing by exactly five edits - and the diff says so.

    A move (a LINE shifted 10 mm), a text change (PUMP A -> PUMP B), an
    attribute change (TAG T101 -> T102), an add (a new LINE) and a delete (a
    LINE removed); a CIRCLE stays untouched. The edits are made on a copy read
    back from the first file, so every surviving entity keeps its handle. The
    diff must list the added line as added, the removed line as removed, the
    three edited entities as changed - and nothing else, the circle included.
    """
    import json

    import ezdxf

    from engineering.understand.diff import diff_snapshots
    from engineering.understand.snapshot import read_snapshot

    folder = tempfile.mkdtemp(prefix="acadmcp_diff_")
    old_path, new_path = os.path.join(folder, "rev_a.dxf"), os.path.join(folder, "rev_b.dxf")
    doc = ezdxf.new("R2018")
    msp = doc.modelspace()
    moved = msp.add_line((0, 0), (100, 0), dxfattribs={"layer": "PIPE"})
    text = msp.add_text("PUMP A", dxfattribs={"insert": (0, 50), "height": 2.5})
    block = doc.blocks.new("TAGGED")
    block.add_circle((0, 0), 10)
    block.add_attdef("TAG", (0, 0), dxfattribs={"height": 2.5})
    insert = msp.add_blockref("TAGGED", (200, 0))
    insert.add_auto_attribs({"TAG": "T101"})
    gone = msp.add_line((0, 100), (100, 100), dxfattribs={"layer": "PIPE"})
    kept = msp.add_circle((300, 300), 20)
    doc.saveas(old_path)

    revision = ezdxf.readfile(old_path)
    revision.entitydb[moved.dxf.handle].translate(10, 0, 0)
    revision.entitydb[text.dxf.handle].dxf.text = "PUMP B"
    revision.entitydb[insert.dxf.handle].get_attrib("TAG").dxf.text = "T102"
    revision.modelspace().delete_entity(revision.entitydb[gone.dxf.handle])
    added = revision.modelspace().add_line((500, 500), (600, 500), dxfattribs={"layer": "PIPE"})
    revision.saveas(new_path)

    diff = diff_snapshots(read_snapshot(old_path), read_snapshot(new_path))
    changed = json.dumps(diff["changed"])
    dumped = changed + json.dumps(diff["added"]) + json.dumps(diff["removed"])
    return (
        len(diff["added"]) == 1
        and len(diff["removed"]) == 1
        and len(diff["changed"]) == 3
        and _quoted(added.dxf.handle) in json.dumps(diff["added"])
        and _quoted(gone.dxf.handle) in json.dumps(diff["removed"])
        and all(_quoted(e.dxf.handle) in changed for e in (moved, text, insert))
        and _quoted(kept.dxf.handle) not in dumped
    )


async def topology_known_defects():
    """Dangling ends, a near miss and a crossing, placed on purpose - and nothing else.

    On layer PIPE: a closed 1000 x 1000 square (no free end, no defect); a LINE
    from (500, -500) that stops at (500, -5), 5 mm short of the square's bottom
    edge (a near miss at gap 10, and a dangling end); two LINEs crossing at
    (2500, 500), away from all four of their ends (one interior crossing); and a
    clean T - a LINE ending on the middle of another at (4500, 0), which is a
    junction, not a crossing and not a near miss.
    """
    import json

    import ezdxf

    from engineering.understand.snapshot import read_snapshot
    from engineering.understand.topology import topology_findings

    doc = ezdxf.new("R2018")
    msp = doc.modelspace()
    pipe = {"layer": "PIPE"}
    ring = msp.add_lwpolyline(
        [(0, 0), (1000, 0), (1000, 1000), (0, 1000)], close=True, dxfattribs=pipe
    )
    short = msp.add_line((500, -500), (500, -5), dxfattribs=pipe)
    cross_a = msp.add_line((2000, 0), (3000, 1000), dxfattribs=pipe)
    cross_b = msp.add_line((2000, 1000), (3000, 0), dxfattribs=pipe)
    bar = msp.add_line((4000, 0), (5000, 0), dxfattribs=pipe)
    stem = msp.add_line((4500, 0), (4500, 500), dxfattribs=pipe)
    path = os.path.join(tempfile.mkdtemp(prefix="acadmcp_topo_"), "topology.dxf")
    doc.saveas(path)

    found = topology_findings(read_snapshot(path), layers=["PIPE"], gap=10.0)
    near = json.dumps(found["near_miss"])
    crossing = json.dumps(found["crossing"])
    dangling = json.dumps(found["dangling"])
    return (
        _quoted(short.dxf.handle) in near
        and _quoted(short.dxf.handle) in dangling
        and _quoted(ring.dxf.handle) not in dangling
        and _quoted(ring.dxf.handle) not in crossing
        and len(found["crossing"]) == 1
        and _quoted(cross_a.dxf.handle) in crossing
        and _quoted(cross_b.dxf.handle) in crossing
        and all(_quoted(e.dxf.handle) not in crossing + near for e in (bar, stem))
    )


# ── Reliability (v1.7.0 W1: the three review reproducers + B07) ─────────────


async def cancel_never_overlaps():
    """A cancelled call's worker still owns the document until it ends (B01).

    The first call's worker blocks on an event; its awaiting task is cancelled
    (what an MCP client's Esc does) and a second call is started at once. The
    second must not run while the first worker is still inside the drawing:
    the trace is A then B, never B first. Version-agnostic: ``_async`` is the
    backend's own lifecycle on every release that has an ezdxf engine.
    """
    import threading

    b = await _b()
    entered, release = threading.Event(), threading.Event()
    trace: list[str] = []

    def slow():
        entered.set()
        release.wait(10)
        trace.append("A:end")

    first = asyncio.ensure_future(b._async(slow))
    if not await asyncio.to_thread(entered.wait, 10):
        return False
    first.cancel()
    try:
        await first
    except asyncio.CancelledError:
        pass
    second = asyncio.ensure_future(b._async(lambda: trace.append("B:run")))
    await asyncio.sleep(0.3)
    overlapped = "B:run" in trace  # B ran while A's worker was still blocked
    release.set()
    await asyncio.wait_for(second, 10)
    return not overlapped and trace == ["A:end", "B:run"]


async def qa_scan_finds_entity_5001():
    """The 5,001st entity is read (B02): the base checks read the first 5,000
    entities of the current space, so an offender beyond them was "clean"."""
    b = await _b()
    msp = b._doc.modelspace()
    for index in range(1, 5_001):
        y = 10.0 * index
        msp.add_line((0.0, y), (100.0, y), dxfattribs={"layer": "GEOMETRY"})
    offender = msp.add_line((-300.0, 0.0), (-200.0, 0.0), dxfattribs={"layer": "CONSTRUCTION"})
    issues = await b.drawing_critique(focus=["construction_left"])
    found = [i for i in issues if i.focus == "construction_left"]
    return len(found) == 1 and found[0].handles == [offender.dxf.handle]


async def export_text_never_formula():
    """Exported text never becomes a spreadsheet formula (B03).

    The parts list goes out as CSV and XLSX with designations that start like
    formulas (``=1+1``, ``@x``, a leading tab) beside a genuine ``-12.5`` mass.
    Read back the way a reader of the file would: ``csv`` gets a leading ``'``
    on each text cell, never on the number; openpyxl finds string cells (data
    type ``s``) for the text and a numeric cell for the mass.
    """
    import csv

    from engineering.sheet.extract import write_rows

    texts = ("=1+1", "@x", "\tD")
    rows = [
        {"item": i, "qty": 1, "designation": t, "material": "S235", "mass": "-12.5"}
        for i, t in enumerate(texts, start=1)
    ]
    columns = ("item", "qty", "designation", "mass")
    with tempfile.TemporaryDirectory() as tmp:
        target = os.path.join(tmp, "bom.csv")
        write_rows(rows, target, columns=columns, fmt="csv")
        with open(target, newline="", encoding="utf-8") as handle:
            body = list(csv.reader(handle))[1:]
        csv_ok = [r[2] for r in body] == ["'" + t for t in texts] and all(
            r[3] == "-12.5" for r in body
        )
        from openpyxl import load_workbook

        target = os.path.join(tmp, "bom.xlsx")
        write_rows(rows, target, columns=columns, fmt="xlsx")
        sheet = load_workbook(target).active
        cells = [(sheet.cell(r, 3), sheet.cell(r, 4)) for r in range(2, 2 + len(texts))]
        xlsx_ok = all(
            text_cell.data_type == "s"
            and text_cell.value == t
            and mass_cell.data_type == "n"
            and mass_cell.value == -12.5
            for (text_cell, mass_cell), t in zip(cells, texts, strict=True)
        )
    return csv_ok and xlsx_ok


async def spec_refuses_non_finite():
    """A non-finite spec value is refused by its path, dry run included (B07).

    A NaN vessel width was accepted (and an inf height died later as a
    RuntimeError), and ``pid_from_spec``'s dry run approved a NaN symbol scale.
    """
    import copy

    from engineering.pid.spec import EXAMPLE_SPEC, run_spec
    from engineering.pid.symbols_vessels import build_vessel

    refused: list[bool] = []
    for params, path in (
        ({"width": float("nan")}, "vertical_vessel: params.width"),
        ({"height": float("inf")}, "vertical_vessel: params.height"),
    ):
        try:
            build_vessel("vertical_vessel", params)
            refused.append(False)
        except Exception as exc:  # noqa: BLE001 - a RuntimeError later on is the old failure
            refused.append(isinstance(exc, ValueError) and str(exc).startswith(path))
    b = await _b()
    spec = copy.deepcopy(EXAMPLE_SPEC)
    spec["equipment"][0]["scale"] = float("nan")
    try:
        await run_spec(b, spec, dry_run=True)
        refused.append(False)
    except Exception as exc:  # noqa: BLE001 - only a refusal by path counts
        refused.append(isinstance(exc, ValueError) and str(exc).startswith("equipment[0].scale"))
    return all(refused)


# ── Verification (v1.7.0 W3, B56: the roadmap's acceptance scenarios) ────────
#
# One check per scenario of the roadmap's acceptance table: T05, T09-T13, T18,
# T25, T36. Each check stages the defect on purpose and passes only when the
# defect is reported AND the clean case is not - a check that cannot fail is
# no evidence. Expected values are worked out by hand in each docstring.


async def _mm_backend():
    """A fresh document declaring millimetres ($INSUNITS 4): verification holds a
    drawing to its declared unit, and a new ezdxf document declares metres."""
    b = await _b()
    await b.system_set_variable("INSUNITS", 4)
    return b


def _requirement_set(*items):
    from engineering.verification.requirements import parse_requirements

    return parse_requirements({"units": "mm", "items": list(items)})


def _req_item(rid, quantity, features, target=None):
    item = {"id": rid, "quantity": quantity, "features": features, "provenance": "explicit"}
    if target is not None:
        item["target"] = target
    return item


def _rows(report) -> dict:
    return {row["id"]: row for row in report["rows"]}


def _plate_with_two_holes():
    """An 80 x 50 plate, Ø10 holes h1 at (40, 25) and h2 at (60, 25): 20 apart."""
    from engineering.mech.part import build_part

    return build_part(
        {
            "kind": "prismatic",
            "name": "plate",
            "material": "aluminium",
            "thickness": 8.0,
            "outline": [[0.0, 0.0], [80.0, 0.0], [80.0, 50.0], [0.0, 50.0]],
            "features": [
                {"kind": "hole", "id": "h1", "x": 40.0, "y": 25.0, "diameter": 10.0},
                {"kind": "hole", "id": "h2", "x": 60.0, "y": 25.0, "diameter": 10.0},
            ],
        }
    )


async def _circle_handle_at(b, x, y):
    for info in await b.entity_list(type_filter="CIRCLE", limit=1000):
        if info.properties["center"][:2] == [x, y]:
            return info.handle
    return None


async def verify_t05_stale_preview_no_silent_mutation():
    """T05: after a preview the operator moves the feature in CAD - stale, and
    nothing of the planned change runs.

    Through the real ``cad_batch`` tool over an in-memory client. The planned
    change moves the line +5 in x. The operator moves its start to (3, 3)
    directly in the document (the content revision cannot see that; the exact
    fingerprint does). The apply must be refused ``stale_revision`` and the line
    must keep the operator's (3, 3) - not (8, 3). A server-side edit between
    preview and apply is refused too. A preview nobody disturbed applies.
    """
    from fastmcp import Client

    import server

    b = await _b()

    async def make():
        return b

    line = await b.entity_create_line(0, 0, 10, 0)
    handle = line.handle
    steps = [{"tool": "entity_move", "args": {"handle": handle, "dx": 5.0, "dy": 0.0}}]

    def start():
        point = b._doc.entitydb[handle].dxf.start
        return (point.x, point.y)

    # the test suite runs this check in-process too: restore the factory after
    original = server._make_backend
    server._make_backend = make
    try:
        return await _t05_through_the_tool(server, Client, b, handle, steps, start)
    finally:
        server._make_backend = original


async def _t05_through_the_tool(server, client_cls, b, handle, steps, start):
    async with client_cls(server.mcp) as client:

        async def preview():
            dry = await client.call_tool("cad_batch", {"steps": steps, "dry_run": True})
            return dry.structured_content["preview"]

        async def apply(planned):
            return await client.call_tool(
                "cad_batch",
                {
                    "steps": steps,
                    "expected_revision": planned["revision"],
                    "preconditions": planned["preconditions"],
                },
                raise_on_error=False,
            )

        planned = await preview()
        b._doc.entitydb[handle].dxf.start = (3.0, 3.0)  # the operator, in CAD
        direct = await apply(planned)
        direct_ok = (
            direct.is_error
            and direct.structured_content.get("kind") == "stale_revision"
            and start() == (3.0, 3.0)
        )
        planned = await preview()
        await client.call_tool("entity_set_properties", {"handle": handle, "color": 1})
        served = await apply(planned)
        served_ok = (
            served.is_error
            and served.structured_content.get("kind") == "stale_revision"
            and start() == (3.0, 3.0)
        )
        clean = await apply(await preview())
        clean_ok = not clean.is_error and start() == (8.0, 3.0)
    return direct_ok and served_ok and clean_ok


async def verify_t09_diameter_under_its_label():
    """T09: a circle drawn Ø39.8 (r = 19.9) under a dimension that reads "Ø40".

    The requirement is Ø40, bound through the dimension: the row is measured on
    the circle (39.8, delta -0.2) and fails; the label's 40 is reported as the
    displayed value and as a ``dimension_mismatch`` finding, and changes
    nothing in the row. The same requirement bound to the circle by handle fails
    with the same 39.8.
    """
    import ezdxf

    from engineering.verification.report import verify_requirements

    doc = ezdxf.new("R2010", setup=True, units=4)
    msp = doc.modelspace()
    circle = msp.add_circle((50, 50), 19.9)
    dim = msp.add_diameter_dim(
        center=(50, 50), radius=19.9, angle=0, text="Ø40", dimstyle="Standard"
    )
    dim.render()
    path = os.path.join(tempfile.mkdtemp(), "t09.dxf")
    doc.saveas(path)
    b = await _b()
    await b.drawing_open(path)
    target = {"nominal": 40.0}
    report = await verify_requirements(
        b,
        _requirement_set(
            _req_item(
                "BY_LABEL",
                "diameter",
                [{"kind": "dimension", "handle": dim.dimension.dxf.handle}],
                target,
            ),
            _req_item(
                "BY_HANDLE", "diameter", [{"kind": "handle", "handle": circle.dxf.handle}], target
            ),
        ),
    )
    rows = _rows(report)
    label, plain = rows["BY_LABEL"], rows["BY_HANDLE"]
    findings = {(f["requirement"], f["kind"]) for f in report["findings"]}
    return (
        label["status"] == "fail"
        and abs(label["measured"]["value"] - 39.8) < 1e-9
        and abs(label["delta"] + 0.2) < 1e-9
        and label["dimension"]["displayed_value"] == 40.0
        and ("BY_LABEL", "dimension_mismatch") in findings
        and plain["status"] == "fail"
        and abs(plain["measured"]["value"] - 39.8) < 1e-9
        and report["accepted"] is False
    )


async def verify_t10_moved_circle_against_its_model():
    """T10: the part model (XDATA) still says h2 is at x = 60; the circle was moved.

    ``draw_part`` stores the model on the drawing and tags each circle with its
    feature. h2 is then moved 2 mm right by ``entity_move`` (the tag moves with
    it). The stored model still reads 60, so it would say the pitch is 20; the
    requirement "h1 to h2 is 20" bound by feature tag measures the circles -
    22 - and fails. The model is not the evidence.
    """
    from engineering.mech.draw import draw_part, read_part
    from engineering.verification.report import verify_requirements

    b = await _mm_backend()
    drawn = await draw_part(b, _plate_with_two_holes(), at=(0.0, 0.0), dimension=False)
    h2 = await _circle_handle_at(b, 60.0, 25.0)
    if h2 is None:
        return False
    await b.entity_move(h2, 2.0, 0.0)
    model = await read_part(b, drawn["part_id"])
    model_x = {
        f.get("id"): f.get("x") for f in (model.get("part") or model).get("features", [])
    }.get("h2")

    def tag(feature):
        return {
            "kind": "feature_tag",
            "part": drawn["part_id"],
            "feature": feature,
            "type": "circle",
        }

    report = await verify_requirements(
        b,
        _requirement_set(_req_item("PITCH", "distance", [tag("h1"), tag("h2")], {"nominal": 20.0})),
    )
    row = _rows(report)["PITCH"]
    return (
        model_x == 60.0
        and row["status"] == "fail"
        and abs(row["measured"]["value"] - 22.0) < 1e-9
        and report["accepted"] is False
    )


async def verify_t11_inner_hole_moved_caught_on_reopen():
    """T11: an inner hole moved inside the outline of the saved file.

    A plate with Ø10 holes at (40, 25) and (60, 25) is saved; in the FILE the
    second hole moves to (62, 25). The entity counts and the extents - what
    1.6's delivery parity compared - are identical, so parity passes. The
    saved file, measured fresh, fails the 20 mm pitch with 22, both re-bound
    from the document and when a fresh engine opens it alone.
    """
    from collections import Counter

    import ezdxf
    from ezdxf import bbox

    from engineering.mech.draw import draw_part
    from engineering.verification.artifact import verify_artifact
    from engineering.verification.report import verify_requirements

    b = await _mm_backend()
    drawn = await draw_part(b, _plate_with_two_holes(), at=(0.0, 0.0), dimension=False)

    def tag(feature):
        return {
            "kind": "feature_tag",
            "part": drawn["part_id"],
            "feature": feature,
            "type": "circle",
        }

    rs = _requirement_set(_req_item("PITCH", "distance", [tag("h1"), tag("h2")], {"nominal": 20.0}))
    if (await verify_requirements(b, rs))["status"] != "pass":
        return False
    path = os.path.join(tempfile.mkdtemp(), "t11.dxf")
    await b.drawing_save_as(path, "dxf")

    def parity(doc):
        msp = doc.modelspace()
        box = bbox.extents(msp)
        return Counter(e.dxftype() for e in msp), (tuple(box.extmin), tuple(box.extmax))

    doc = ezdxf.readfile(path)
    before = parity(doc)
    moved = 0
    for circle in doc.modelspace().query("CIRCLE"):
        if tuple(circle.dxf.center)[:2] == (60.0, 25.0):
            circle.dxf.center = (62.0, 25.0, 0.0)
            moved += 1
    doc.saveas(path)
    parity_blind = parity(ezdxf.readfile(path)) == before
    on_file = _rows(await verify_artifact(path, rs, backend=b))["PITCH"]
    fresh = await _b()
    await fresh.drawing_open(path)
    alone = _rows(await verify_requirements(fresh, rs))["PITCH"]
    return (
        moved == 1
        and parity_blind
        and on_file["status"] == "fail"
        and abs(on_file["measured"]["value"] - 22.0) < 1e-9
        and alone["status"] == "fail"
    )


async def verify_t12_equal_holes_ambiguous():
    """T12: two equal Ø10 holes and a requirement that names "the" Ø10 hole.

    The binding must be ``ambiguous`` with both holes as candidates and no
    handle chosen; the report is not a pass and is not accepted. Naming the
    hole by its feature tag binds exactly one.
    """
    from engineering.mech.draw import draw_part
    from engineering.verification.report import verify_requirements

    b = await _mm_backend()
    drawn = await draw_part(b, _plate_with_two_holes(), at=(0.0, 0.0), dimension=False)
    window = {"kind": "predicate", "type": "circle", "radius_min": 4.999, "radius_max": 5.001}
    tagged = {"kind": "feature_tag", "part": drawn["part_id"], "feature": "h1", "type": "circle"}
    report = await verify_requirements(
        b,
        _requirement_set(
            _req_item("ANY", "diameter", [window], {"nominal": 10.0}),
            _req_item("H1", "diameter", [tagged], {"nominal": 10.0}),
        ),
    )
    rows = _rows(report)
    expected = sorted(
        [await _circle_handle_at(b, 40.0, 25.0), await _circle_handle_at(b, 60.0, 25.0)]
    )
    candidates = sorted(
        handle for ref in rows["ANY"]["binding"]["candidates"] for handle in ref["handles"]
    )
    return (
        rows["ANY"]["status"] == "ambiguous"
        and candidates == expected
        and not rows["ANY"]["handles"]
        and rows["H1"]["status"] == "pass"
        and report["status"] != "pass"
        and report["accepted"] is False
    )


async def verify_t13_hole_count_cardinality():
    """T13: exactly six Ø8 holes on a Ø100 PCD - five and seven both fail.

    Every hole has the right diameter in all three drawings, so a check that
    only measured diameters would pass them all. The count requirement passes
    six, fails five (one deleted) and fails seven (one added).
    """
    from engineering.verification.report import verify_requirements

    b = await _mm_backend()
    handles = []
    for k in range(6):
        angle = math.radians(60.0 * k)
        c = await b.entity_create_circle(50 * math.cos(angle), 50 * math.sin(angle), 4.0)
        handles.append(c.handle)
    window = {"kind": "predicate", "type": "circle", "radius_min": 3.999, "radius_max": 4.001}
    rs = _requirement_set(_req_item("N", "count", [window], {"exact": 6}))

    async def count_status():
        row = _rows(await verify_requirements(b, rs))["N"]
        return row["status"], row["measured"]["value"]

    six = await count_status()
    await b.entity_delete(handles[0])
    five = await count_status()
    await b.entity_create_circle(50.0, 0.0, 4.0)
    await b.entity_create_circle(0.0, 0.0, 4.0)
    seven = await count_status()
    return six == ("pass", 6.0) and five == ("fail", 5.0) and seven == ("fail", 7.0)


async def verify_t18_dimension_override_and_dimlfac_visible():
    """T18: the geometry and the annotation are judged apart.

    Two vertical lines 40 apart under a dimension whose text is overridden to
    "42", and two lines 21 apart under a dimension with DIMLFAC 2 (it shows
    42). Each requirement matches its geometry (40, 21), so each row passes -
    and each dimension is a ``dimension_mismatch`` finding that blocks
    acceptance. The same lines under an honest dimension are accepted.
    """
    import ezdxf

    from engineering.verification.report import verify_requirements

    def sheet(**kw):
        doc = ezdxf.new("R2010", setup=True, units=4)
        msp = doc.modelspace()
        handles = []
        for x0, width, extra in (
            (0.0, 40.0, kw.get("first", {})),
            (100.0, 21.0, kw.get("second", {})),
        ):
            msp.add_line((x0, 0), (x0, 20))
            msp.add_line((x0 + width, 0), (x0 + width, 20))
            dim = msp.add_linear_dim(
                base=(x0, 30), p1=(x0, 0), p2=(x0 + width, 0), dimstyle="Standard", **extra
            )
            dim.render()
            handles.append(dim.dimension.dxf.handle)
        path = os.path.join(tempfile.mkdtemp(), "t18.dxf")
        doc.saveas(path)
        return path, handles

    async def measure(path, handles):
        b = await _b()
        await b.drawing_open(path)
        return await verify_requirements(
            b,
            _requirement_set(
                _req_item(
                    "W40",
                    "distance",
                    [{"kind": "dimension", "handle": handles[0]}],
                    {"nominal": 40.0},
                ),
                _req_item(
                    "W21",
                    "distance",
                    [{"kind": "dimension", "handle": handles[1]}],
                    {"nominal": 21.0},
                ),
            ),
        )

    lying = await measure(*sheet(first={"text": "42"}, second={"override": {"dimlfac": 2.0}}))
    honest = await measure(*sheet())
    rows = _rows(lying)
    kinds = sorted((f["requirement"], f["kind"]) for f in lying["findings"])
    return (
        rows["W40"]["status"] == "pass"
        and rows["W21"]["status"] == "pass"
        and rows["W40"]["dimension"]["displayed_value"] == 42.0
        and rows["W21"]["dimension"]["displayed_value"] == 42.0
        and kinds == [("W21", "dimension_mismatch"), ("W40", "dimension_mismatch")]
        and lying["accepted"] is False
        and honest["accepted"] is True
        and honest["findings"] == []
    )


async def verify_t25_handle_renumber_rebinds():
    """T25: a save that renumbers handles - re-bound by geometry, never by a
    stale handle; a copied feature tag is ``unknown``, never a pick.

    Two files hold the same Ø40 bore (r = 20) and a Ø10 pin (r = 5), created in
    opposite orders, so the saved file gave the bore's old handle to the pin.
    Re-bound from the source, the bore still measures 40 and passes. When the
    saved file also holds a copy carrying the bore's feature tag, the binding
    has two candidates and the row is ``unknown``.
    """
    import ezdxf

    from engineering.mech.xdata import encode
    from engineering.verification.artifact import verify_artifact

    tag = {"kind": "feature", "part": "P1", "feature": "bore", "view": "front"}
    folder = tempfile.mkdtemp()

    def write(name, circles):
        doc = ezdxf.new("R2010", units=4)
        doc.appids.add("ACADMCP_MECH")
        msp = doc.modelspace()
        msp.add_lwpolyline([(0, 0), (200, 0), (200, 100), (0, 100)], close=True)
        handles = {}
        for key, cx, cy, radius, tagged in circles:
            circle = msp.add_circle((cx, cy), radius)
            if tagged:
                rows = encode(tag, kinds=("feature",))
                circle.set_xdata("ACADMCP_MECH", [row for row in rows if row[0] != 1001])
            handles[key] = circle.dxf.handle
        path = os.path.join(folder, name)
        doc.saveas(path)
        return path, handles

    source, before = write("a.dxf", [("bore", 50, 50, 20, True), ("pin", 150, 80, 5, False)])
    renumbered, after = write("b.dxf", [("pin", 150, 80, 5, False), ("bore", 50, 50, 20, True)])
    copied, _ = write(
        "c.dxf",
        [("copy", 150, 50, 20, True), ("bore", 50, 50, 20, True), ("pin", 150, 80, 5, False)],
    )
    if after["pin"] != before["bore"]:
        return False  # the fixture must renumber, or it proves nothing
    b = await _b()
    await b.drawing_open(source)
    rs = _requirement_set(
        _req_item(
            "BORE", "diameter", [{"kind": "handle", "handle": before["bore"]}], {"nominal": 40.0}
        )
    )
    rebound = _rows(await verify_artifact(renumbered, rs, backend=b))["BORE"]
    doubled = _rows(await verify_artifact(copied, rs, backend=b))["BORE"]
    return (
        rebound["status"] == "pass"
        and abs(rebound["measured"]["value"] - 40.0) < 1e-9
        and rebound["handles"] == [after["bore"]]
        and doubled["status"] == "unknown"
        and len(doubled["binding"]["candidates"]) == 2
    )


async def verify_t36_report_stale_after_file_change():
    """T36: the delivered file changes after its report was written.

    A Ø40 bore is delivered with its requirement: the manifest re-hashes clean
    and the acceptance report is accepted. One line is then appended to the
    delivered DXF. ``verify_manifest`` must name the DXF as stale, and the
    acceptance report, re-checked against its artifact's sha256, must turn
    stale and stop being accepted.
    """
    import json

    import ezdxf

    from engineering.delivery import deliver_drawing, verify_manifest
    from engineering.layers import ensure_engineering_layers, ensure_standard_linetypes
    from engineering.verification.report import mark_stale

    b = await _mm_backend()
    await ensure_standard_linetypes(b)
    await ensure_engineering_layers(b)
    bore = await b.entity_create_circle(50, 50, 20, layer="GEOMETRY")
    rs = _requirement_set(
        _req_item(
            "BORE", "diameter", [{"kind": "handle", "handle": bore.handle}], {"nominal": 40.0}
        )
    )
    bundle = os.path.join(tempfile.mkdtemp(), "bundle")
    result = await deliver_drawing(
        b, bundle, formats=["dxf"], min_score=0, strict_critique=False, requirements=rs
    )
    manifest = os.path.join(bundle, "manifest.json")
    clean = verify_manifest(manifest)
    with open(os.path.join(bundle, "acceptance.json"), encoding="utf-8") as handle:
        report = json.load(handle)
    accepted_before = mark_stale(report, None)["accepted"]
    delivered = os.path.join(bundle, "drawing.dxf")
    doc = ezdxf.readfile(delivered)
    doc.modelspace().add_line((0, 0), (10, 0))
    doc.saveas(delivered)
    edited = verify_manifest(manifest)
    later = mark_stale(report, None)
    return (
        result.status == "success"
        and clean["ok"] is True
        and accepted_before is True
        and edited["ok"] is False
        and [row["filename"] for row in edited["stale"]] == ["drawing.dxf"]
        and later["stale"] is True
        and later["accepted"] is False
    )


CHECKS = {
    "core_line_length": (core_line_length, "Core"),
    "core_circle_radius": (core_circle_radius, "Core"),
    "core_layer_create": (core_layer_create, "Core"),
    "core_polyline_closed": (core_polyline_closed, "Core"),
    "core_linear_dim": (core_linear_dim, "Core"),
    "core_save_dxf_roundtrip": (core_save_dxf_roundtrip, "Core"),
    "dim_aligned_no_error": (dim_aligned_no_error, "Dimensions"),
    "dim_angular_no_error": (dim_angular_no_error, "Dimensions"),
    "array_polar_360_distinct": (array_polar_360_distinct, "Modify"),
    "point_intersection_line_line": (point_intersection_line_line, "Geometry"),
    "point_tangent_external": (point_tangent_external, "Geometry"),
    "arc_has_length": (arc_has_length, "Query"),
    "arc_select_by_length": (arc_select_by_length, "Query"),
    "ezdxf_bounding_box": (ezdxf_bounding_box, "Query"),
    "mtext_rotation_roundtrip": (mtext_rotation_roundtrip, "Entities"),
    "screenshot_png": (screenshot_png, "Render"),
    "dim_overlap_critique_fires": (dim_overlap_critique_fires, "Quality gate"),
    "iso13567_dim_layer": (iso13567_dim_layer, "Quality gate"),
    "construction_left_iso_caught": (construction_left_iso_caught, "Quality gate"),
    "gear_no_self_overlap": (gear_no_self_overlap, "Engineering"),
    "center_linetype_applied": (center_linetype_applied, "Entities"),
    "hatch_area_subtracts_its_island": (hatch_area_subtracts_its_island, "Measurement"),
    "boundary_area_agrees_with_measure": (boundary_area_agrees_with_measure, "Measurement"),
    "two_vertex_circle_has_area": (two_vertex_circle_has_area, "Measurement"),
    "diameter_dim_measures_the_diameter": (diameter_dim_measures_the_diameter, "Dimensions"),
    "radius_dim_ignores_the_leader_length": (radius_dim_ignores_the_leader_length, "Dimensions"),
    "pid_block_define_attdef_roundtrip": (pid_block_define_attdef_roundtrip, "P&ID"),
    "pid_tag_parse_fic": (pid_tag_parse_fic, "P&ID"),
    "pid_graph_edge_count": (pid_graph_edge_count, "P&ID"),
    "settings_dimstyle_iso25_values": (settings_dimstyle_iso25_values, "Settings"),
    "settings_layer_state_roundtrip": (settings_layer_state_roundtrip, "Settings"),
    "settings_pdf_mediabox_a3": (settings_pdf_mediabox_a3, "Settings"),
    "mech_part_roundtrip": (mech_part_roundtrip, "Mechanical"),
    "mech_section_hatch_area": (mech_section_hatch_area, "Mechanical"),
    "mech_iso286_on_dimension": (mech_iso286_on_dimension, "Mechanical"),
    "mech_thread_unrepresented_is_caught": (mech_thread_unrepresented_is_caught, "Mechanical"),
    "std_part_iso4014_m12": (std_part_iso4014_m12, "Mechanical"),
    "sheet_frame_iso5457_a3": (sheet_frame_iso5457_a3, "Sheet"),
    "bom_balloon_link": (bom_balloon_link, "Sheet"),
    "arch_junction_l_t_x": (arch_junction_l_t_x, "Architecture"),
    "arch_room_area_net": (arch_room_area_net, "Architecture"),
    "arch_opening_cuts_wall": (arch_opening_cuts_wall, "Architecture"),
    "arch_rooms_detect_foreign": (arch_rooms_detect_foreign, "Architecture"),
    "scale_check_detects_schematic": (scale_check_detects_schematic, "Understanding"),
    "pipe_takeoff_rmst_exact": (pipe_takeoff_rmst_exact, "Understanding"),
    "cable_takeoff_roundup": (cable_takeoff_roundup, "Understanding"),
    "diff_detects_known_edits": (diff_detects_known_edits, "Understanding"),
    "topology_known_defects": (topology_known_defects, "Understanding"),
    "cancel_never_overlaps": (cancel_never_overlaps, "Reliability"),
    "qa_scan_finds_entity_5001": (qa_scan_finds_entity_5001, "Reliability"),
    "export_text_never_formula": (export_text_never_formula, "Reliability"),
    "spec_refuses_non_finite": (spec_refuses_non_finite, "Reliability"),
    "verify_t05_stale_preview_no_silent_mutation": (
        verify_t05_stale_preview_no_silent_mutation,
        "Verification",
    ),
    "verify_t09_diameter_under_its_label": (verify_t09_diameter_under_its_label, "Verification"),
    "verify_t10_moved_circle_against_its_model": (
        verify_t10_moved_circle_against_its_model,
        "Verification",
    ),
    "verify_t11_inner_hole_moved_caught_on_reopen": (
        verify_t11_inner_hole_moved_caught_on_reopen,
        "Verification",
    ),
    "verify_t12_equal_holes_ambiguous": (verify_t12_equal_holes_ambiguous, "Verification"),
    "verify_t13_hole_count_cardinality": (verify_t13_hole_count_cardinality, "Verification"),
    "verify_t18_dimension_override_and_dimlfac_visible": (
        verify_t18_dimension_override_and_dimlfac_visible,
        "Verification",
    ),
    "verify_t25_handle_renumber_rebinds": (verify_t25_handle_renumber_rebinds, "Verification"),
    "verify_t36_report_stale_after_file_change": (
        verify_t36_report_stale_after_file_change,
        "Verification",
    ),
}


def main():
    if len(sys.argv) == 1 or sys.argv[1] in {"-h", "--help"}:
        print("usage: correctness_suite.py <check_name> | --list")
        print("Run one deterministic correctness check, or list available checks.")
        return
    if len(sys.argv) == 2 and sys.argv[1] == "--list":
        for k, (_fn, cat) in CHECKS.items():
            print(f"{k}\t{cat}")
        return
    name = sys.argv[1]
    if name not in CHECKS:
        print(f"Unknown check: {name}", file=sys.stderr)
        raise SystemExit(2)
    fn, _cat = CHECKS[name]
    ok = asyncio.run(fn())
    print("PASS" if ok else "FAIL")


if __name__ == "__main__":
    main()
